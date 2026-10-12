"""make_pptx.py: the Vellum slides as a PowerPoint template, light and dark, that Canva can import.

Writes docs/downloads/vellum-slides-light.pptx and vellum-slides-dark.pptx (16:9, 13.333 x 7.5 in). Each file is a
template, not a set of pictures:

- the theme (ppt/theme/theme1.xml) is rewritten: colour scheme "Vellum" with dk1 ink, lt1 surface, dk2 secondary, lt2
  card, accent1-6 the six tones in slot order (teal, terracotta, indigo, ochre, mauve, moss), hyperlink indigo; font
  scheme "Vellum" with Palatino Linotype as the major (heading) font and Corbel as the minor (body) font. The dark file
  keeps the slots by role (dk1 is still the text colour, now the light ink of the dark twin), so the default colour
  map holds and no remapping is needed in PowerPoint or Canva.
- the slide master is rewritten: background bg1, the teal heading rule, the title, body, footer and slide-number
  placeholders, and the text styles (claim 32 pt serif in ink, body 20 pt in secondary).
- ten named layouts with placeholders: Title, Section, Claim + figure, Claim + two figures, Big number, Chart, Table,
  Claim + text, Claim only, Closing.
- thirteen example slides use them with real Vellum content: gallery PNGs (the dark deck takes the *_dark twins,
  cropped of their own title block and footer, since the slide carries the claim and the source), two native charts
  (python-pptx) styled the Vellum way, a big number and stat tiles on card-coloured rounded rectangles, a native table,
  and a last slide on using the template in PowerPoint and Canva.

Type sizes follow vellum.py's slide scale (2.2): claim 32 pt, body 20 pt, figure and chart text 16 pt, footer 11 pt.
Text is in ink, secondary or muted only; the tones are for data marks.

Then, when PowerPoint is installed, the decks are opened read-only through COM (no window; an open that needs a
repair fails here instead of prompting), every slide is exported at 1600 x 900 px to docs/media/slides/<mode>/
slide-NN.png, the light deck is saved as PDF next to them, and the deck is recorded in docs/figures/slides_pptx/.
Run: python examples/python/office/make_pptx.py [--no-render]
"""
import copy, csv, io, os, subprocess, sys, tempfile

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, ".."))
sys.path.insert(0, os.path.join(HERE, "..", "infographic"))
from _common import DATA, DOCS, FIGURES, record, rel, stamp, vellum          # noqa: E402
from story_kit import facts, num, SOURCE, DATE, AUTHOR, URL, REGIONS        # noqa: E402

SPEC = vellum.SPEC
DOWNLOADS = os.path.join(DOCS, "downloads")
MEDIA = os.path.join(DOCS, "media", "slides")
SCRIPT = os.path.basename(__file__)
GIT = stamp()

# ---- the grid, in inches, on a 13.333 x 7.5 in slide
SW, SH = 13.333, 7.5
M = 0.667                                  # side margin: 8 mm of the 160 mm beamer frame, as in the pitch deck
CW = SW - 2 * M                            # content width, 12.0 in
RULE_Y, RULE_W, RULE_H = 0.50, 0.533, 0.042
TITLE_Y = 0.66                             # top of the claim
FOOT_Y, FOOT_H = 6.98, 0.28                # the footer line
CONTENT_BOTTOM = 6.72                      # nothing above the footer goes lower
GUTTER = 0.4
BIG_Y = 2.0                                # top of the big-number tile

# ---- type, in points
CLAIM, BODY, UNIT, FIG, FOOT = 32, 20, 18, 16, 11
COVER, SECTION, HERO, STAT = 44, 40, 88, 60

LINE_PITCH = {"title": 1.184, "body": 1.22}
FONT_FILES = {"title": "C:/Windows/Fonts/pala.ttf", "body": "C:/Windows/Fonts/corbel.ttf"}
FACES = {"title": SPEC["typography"]["title"], "body": SPEC["typography"]["body"]}


# ============================================================================ text measure
_fonts = {}


def _font(face, pt):
    key = (face, pt)
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype(FONT_FILES[face], size=pt * 20)
    return _fonts[key]


def text_w(s, face, pt):
    """Advance width of a string in inches, measured with the face the slide uses."""
    return _font(face, pt).getlength(s) / 20 / 72


def n_lines(s, face, pt, width):
    """Lines a greedy wrap needs for `s` at `pt` in a box `width` inches wide (as PowerPoint wraps, word by word)."""
    lines, cur = 1, ""
    for word in s.split(" "):
        trial = (cur + " " + word).strip()
        if text_w(trial, face, pt) <= width or not cur:
            cur = trial
        else:
            lines, cur = lines + 1, word
    return lines


def line_h(face, pt, spacing=1.0):
    """Height of one line in inches as PowerPoint sets it, times the line spacing. The factors (line pitch per em at
    100 %) are measured on PowerPoint's own exports of this deck."""
    return pt * LINE_PITCH[face] * spacing / 72


def fits(s, face, pt, width, max_lines):
    k = n_lines(s, face, pt, width)
    assert k <= max_lines, f"{s!r} needs {k} lines at {pt} pt in {width:.2f} in; shorten it"
    return k


# ============================================================================ palette
class Palette:
    def __init__(self, mode):
        self.mode = mode
        P = SPEC[mode]; N = P["neutrals"]
        self.surface, self.card, self.grid, self.rule = N["surface"], N["card"], N["grid"], N["rule"]
        self.muted, self.secondary, self.ink = N["muted"], N["secondary"], N["ink"]
        self.tone, self.tint, self.deep = P["tone"], P["tint"], P["deep"]
        self.names = SPEC["order"]

    def t(self, name):
        return self.tone[self.names.index(name)]


def h(c):
    return c.lstrip("#").upper()


def rgb(c):
    return RGBColor.from_string(h(c))


# ============================================================================ theme
def set_theme(prs, V):
    """Colour scheme and font scheme of ppt/theme/theme1.xml, edited in place."""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    root.set("name", "Vellum" if V.mode == "light" else "Vellum dark")
    te = root.find(a + "themeElements")
    slots = [("dk1", V.ink), ("lt1", V.surface), ("dk2", V.secondary), ("lt2", V.card)]
    slots += [(f"accent{k + 1}", V.tone[k]) for k in range(6)]
    slots += [("hlink", V.t("indigo")), ("folHlink", V.deep[SPEC["order"].index("indigo")])]
    cs = etree.fromstring(f'<a:clrScheme {nsdecls("a")} name="Vellum{"" if V.mode == "light" else " dark"}">'
                          + "".join(f'<a:{k}><a:srgbClr val="{h(v)}"/></a:{k}>' for k, v in slots) + "</a:clrScheme>")
    te.replace(te.find(a + "clrScheme"), cs)
    fs = etree.fromstring(
        f'<a:fontScheme {nsdecls("a")} name="Vellum">'
        f'<a:majorFont><a:latin typeface="{FACES["title"]}"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont>'
        f'<a:minorFont><a:latin typeface="{FACES["body"]}"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont>'
        "</a:fontScheme>")
    te.replace(te.find(a + "fontScheme"), fs)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ============================================================================ shape XML for the master and layouts
def E(inches):
    return int(round(inches * 914400))


def xfrm(x, y, w, hh):
    return f'<a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(w)}" cy="{E(hh)}"/></a:xfrm>'


def fill(color=None, scheme=None):
    if scheme: return f'<a:solidFill><a:schemeClr val="{scheme}"/></a:solidFill>'
    return f'<a:solidFill><a:srgbClr val="{h(color)}"/></a:solidFill>'


def rpr(pt, color=None, scheme=None, face="+mn-lt", bold=False):
    b = ' b="1"' if bold else ' b="0"'
    return f'<a:defRPr sz="{int(pt * 100)}"{b} kern="0">{fill(color, scheme)}<a:latin typeface="{face}"/></a:defRPr>'


def ppr(lvl, pt, color=None, scheme=None, face="+mn-lt", algn="l", spacing=100, after=0, bold=False):
    return (f'<a:lvl{lvl}pPr marL="0" indent="0" algn="{algn}"><a:lnSpc><a:spcPct val="{spacing * 1000}"/></a:lnSpc>'
            f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="{int(after * 100)}"/></a:spcAft>'
            f'<a:buNone/>{rpr(pt, color, scheme, face, bold)}</a:lvl{lvl}pPr>')


BODYPR = '<a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}" anchorCtr="0"><a:noAutofit/></a:bodyPr>'


class Shapes:
    """Collects the <p:sp> of a master or layout with ids counted from 2."""
    def __init__(self):
        self.items, self.n = [], 1

    def _id(self):
        self.n += 1
        return self.n

    def ph(self, kind, idx, name, box, prompt="", style="", anchor="t", sz=None):
        i = self._id()
        t = f' type="{kind}"' if kind else ""
        ix = f' idx="{idx}"' if idx is not None else ""
        s = f' sz="{sz}"' if sz else ""
        if kind == "sldNum":
            txt = (f'<p:txBody>{BODYPR.format(anchor=anchor)}<a:lstStyle>{style}</a:lstStyle><a:p>'
                   f'<a:fld id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum"><a:rPr lang="en-GB"/>'
                   f'<a:t>{prompt}</a:t></a:fld></a:p></p:txBody>')
        elif kind not in ("pic", "chart", "tbl"):
            txt = (f'<p:txBody>{BODYPR.format(anchor=anchor)}<a:lstStyle>{style}</a:lstStyle>'
                   f'<a:p><a:r><a:rPr lang="en-GB" dirty="0"/><a:t>{prompt}</a:t></a:r></a:p></p:txBody>')
        else:
            txt = (f'<p:txBody>{BODYPR.format(anchor="t")}<a:lstStyle/><a:p><a:r><a:rPr lang="en-GB" dirty="0"/>'
                   f'<a:t>{prompt}</a:t></a:r></a:p></p:txBody>')
        self.items.append(
            f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph{t}{s}{ix}/></p:nvPr></p:nvSpPr><p:spPr>{xfrm(*box)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>{txt}</p:sp>')

    def rect(self, name, box, scheme="accent1", rounded=None):
        i = self._id()
        geom = (f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val {rounded}"/></a:avLst></a:prstGeom>'
                if rounded is not None else '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>')
        self.items.append(
            f'<p:sp><p:nvSpPr><p:cNvPr id="{i}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
            f'<p:spPr>{xfrm(*box)}{geom}{fill(scheme=scheme)}<a:ln><a:noFill/></a:ln></p:spPr>'
            f'<p:txBody><a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle/><a:p><a:endParaRPr lang="en-GB"/></a:p></p:txBody></p:sp>')

    def tree(self):
        return (f'<p:spTree {nsdecls("p", "a", "r")}><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
                '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
                + "".join(self.items) + "</p:spTree>")


# the boxes every claim layout shares (the defaults for a two-line claim; slides move them for a one-line claim)
TITLE_BOX = (M, TITLE_Y, CW, 1.10)
UNIT_BOX = (M, 1.84, CW, 0.36)
CONTENT_Y = 2.42
FOOT_BOX = (M, FOOT_Y, 10.0, FOOT_H)
NUM_BOX = (M + 10.2, FOOT_Y, CW - 10.2, FOOT_H)


def foot_style(V, algn="l"):
    return ppr(1, FOOT, V.muted, algn=algn)


def common(S, V, unit=True, unit_prompt="Unit or subtitle line: what is shown, in which unit"):
    S.ph("title", None, "Claim", TITLE_BOX, "The claim, in one sentence, as the title")
    if unit:
        S.ph("body", 13, "Unit line", UNIT_BOX, unit_prompt, ppr(1, UNIT, V.muted), sz="quarter")


def footers(S, V):
    S.ph("ftr", 11, "Footer", FOOT_BOX, "Source · script · git stamp", foot_style(V), anchor="b", sz="quarter")
    S.ph("sldNum", 12, "Slide number", NUM_BOX, "‹#›", foot_style(V, "r"), anchor="b", sz="quarter")


def layouts(V):
    """name, layout type (or None for a custom layout), show master shapes, background scheme colour, spTree."""
    out = []
    # 1 Title: the rule above a large serif claim, a subtitle, a muted line for author and date
    S = Shapes()
    S.rect("Rule", (M, 2.06, RULE_W, RULE_H))
    S.ph("ctrTitle", None, "Title", (M, 2.24, 11.4, 1.66), "The talk’s claim, in one sentence",
         ppr(1, COVER, scheme="tx1", face="+mj-lt", spacing=95))
    S.ph("subTitle", 1, "Subtitle", (M, 4.16, 11.4, 0.80), "What the talk shows, in one line",
         ppr(1, 22, scheme="tx2"))
    S.ph("body", 13, "Author and date", (M, 5.20, 11.4, 0.40), "Author · venue · date", ppr(1, FIG, V.muted), sz="quarter")
    footers(S, V)
    out.append(("Title", "title", False, None, S))
    # 2 Section: on the card, a muted number, the rule, the section title, one line
    S = Shapes()
    S.ph("body", 13, "Section number", (M, 2.30, 4.0, 0.40), "01", ppr(1, BODY, V.muted), sz="quarter")
    S.rect("Rule", (M, 2.92, RULE_W, RULE_H))
    S.ph("title", None, "Section title", (M, 3.12, 11.4, 1.10), "Section title",
         ppr(1, SECTION, scheme="tx1", face="+mj-lt", spacing=95))
    S.ph("body", 14, "Section line", (M, 4.40, 11.4, 0.80), "What this part shows, in one line", ppr(1, BODY, scheme="tx2"),
         sz="quarter")
    footers(S, V)
    out.append(("Section", "secHead", False, "bg2", S))
    # 3 Claim + figure
    S = Shapes(); common(S, V)
    S.ph("pic", 14, "Figure", (M, CONTENT_Y, CW, CONTENT_BOTTOM - CONTENT_Y), "Figure: a Vellum PNG at slide scale")
    footers(S, V)
    out.append(("Claim + figure", None, True, None, S))
    # 4 Claim + two figures, each with a caption line
    S = Shapes(); common(S, V)
    half = (CW - GUTTER) / 2; ph_h = CONTENT_BOTTOM - CONTENT_Y - 0.62
    for k, side in enumerate(("left", "right")):
        x = M + k * (half + GUTTER)
        S.ph("pic", 14 + k, f"Figure {side}", (x, CONTENT_Y, half, ph_h), f"Figure, {side}")
        S.ph("body", 16 + k, f"Caption {side}", (x, CONTENT_BOTTOM - 0.46, half, 0.46), f"What the {side} figure shows",
             ppr(1, FIG, scheme="tx2"), sz="quarter")
    footers(S, V)
    out.append(("Claim + two figures", None, True, None, S))
    # 5 Big number: the number on a card tile with its label and comparison; one sentence beside it
    S = Shapes(); common(S, V, unit=False)
    S.rect("Tile", (M, BIG_Y, 6.6, 3.70), scheme="bg2", rounded=3200)
    S.ph("body", 13, "Big number", (M + 0.45, BIG_Y + 0.32, 5.7, 1.55), "−1.7 years",
         ppr(1, HERO, scheme="tx1", face="+mj-lt", spacing=90), anchor="b", sz="quarter")
    S.ph("body", 14, "What the number is", (M + 0.45, BIG_Y + 2.06, 5.7, 0.72), "what the number counts, where and when",
         ppr(1, BODY, scheme="tx2", spacing=95), sz="quarter")
    S.ph("body", 15, "Comparison", (M + 0.45, BIG_Y + 2.82, 5.7, 0.70), "the comparison that gives it scale",
         ppr(1, UNIT, V.muted, spacing=95), sz="quarter")
    S.ph("body", 16, "Sentence", (M + 7.2, BIG_Y + 0.32, CW - 7.2, 3.2), "One sentence on what the number means.",
         ppr(1, 24, scheme="tx1", spacing=105), sz="half")
    footers(S, V)
    out.append(("Big number", None, True, None, S))
    # 6 Chart: a native chart under the unit line
    S = Shapes(); common(S, V)
    S.ph("chart", 14, "Chart", (M, CONTENT_Y, CW, CONTENT_BOTTOM - CONTENT_Y), "Chart: insert a native chart")
    footers(S, V)
    out.append(("Chart", None, True, None, S))
    # 7 Table
    S = Shapes(); common(S, V)
    S.ph("tbl", 14, "Table", (M, CONTENT_Y, CW, CONTENT_BOTTOM - CONTENT_Y), "Table")
    footers(S, V)
    out.append(("Table", None, True, None, S))
    # 8 Claim + text: plain paragraphs, no bullets
    S = Shapes(); common(S, V, unit=False)
    S.ph("body", 1, "Text", (M, 1.98, 11.0, CONTENT_BOTTOM - 1.98), "Paragraphs of 20 pt text in secondary, no bullets",
         ppr(1, BODY, scheme="tx2", spacing=105, after=12))
    footers(S, V)
    out.append(("Claim + text", "obj", True, None, S))
    # 9 Claim only: the claim, the unit line and free space (stat tiles, diagrams)
    S = Shapes(); common(S, V)
    footers(S, V)
    out.append(("Claim only", "titleOnly", True, None, S))
    # 10 Closing: the result in one sentence, the contact, a muted line on where everything comes from
    S = Shapes()
    S.rect("Rule", (M, 1.86, RULE_W, RULE_H))
    S.ph("title", None, "Closing claim", (M, 2.08, 11.4, 1.60), "The result, in one sentence",
         ppr(1, 36, scheme="tx1", face="+mj-lt", spacing=95))
    S.ph("body", 13, "Contact", (M, 4.10, 11.4, 0.45), "Name · web address", ppr(1, BODY, scheme="tx1"), sz="quarter")
    S.ph("body", 14, "Provenance", (M, 4.75, 11.4, 1.10), "Where the figures, data and code are",
         ppr(1, FIG, V.muted, spacing=105), sz="quarter")
    footers(S, V)
    out.append(("Closing", None, False, None, S))
    return out


def master_xml(V):
    S = Shapes()
    S.rect("Rule", (M, RULE_Y, RULE_W, RULE_H))
    S.ph("title", None, "Title", TITLE_BOX, "The claim, in one sentence, as the title")
    S.ph("body", 1, "Text", (M, CONTENT_Y, CW, CONTENT_BOTTOM - CONTENT_Y), "Text in 20 pt Corbel")
    S.ph("dt", 10, "Date", (M + 6.0, FOOT_Y, 3.0, FOOT_H), "", foot_style(V), anchor="b", sz="half")
    S.ph("ftr", 11, "Footer", FOOT_BOX, "", foot_style(V), anchor="b", sz="quarter")
    S.ph("sldNum", 12, "Slide number", NUM_BOX, "‹#›", foot_style(V, "r"), anchor="b", sz="quarter")
    return S.tree()


def tx_styles(V):
    def lvl(n, pt, scheme="tx2", mar=0.0, bullet=False):
        bu = ('<a:buFont typeface="+mj-lt"/><a:buChar char="–"/>' if bullet else "<a:buNone/>")
        ind = -E(0.25) if bullet else 0
        return (f'<a:lvl{n}pPr marL="{E(mar)}" indent="{ind}" algn="l" defTabSz="914400"><a:lnSpc><a:spcPct val="105000"/></a:lnSpc>'
                f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="1000"/></a:spcAft>{bu}'
                f'<a:defRPr sz="{pt * 100}" kern="0">{fill(scheme=scheme)}<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/>'
                f'<a:cs typeface="+mn-cs"/></a:defRPr></a:lvl{n}pPr>')
    title = ('<p:titleStyle><a:lvl1pPr algn="l" defTabSz="914400"><a:lnSpc><a:spcPct val="95000"/></a:lnSpc>'
             f'<a:spcBef><a:spcPct val="0"/></a:spcBef><a:buNone/><a:defRPr sz="{CLAIM * 100}" kern="0">{fill(scheme="tx1")}'
             '<a:latin typeface="+mj-lt"/><a:ea typeface="+mj-ea"/><a:cs typeface="+mj-cs"/></a:defRPr></a:lvl1pPr></p:titleStyle>')
    body = ("<p:bodyStyle>" + lvl(1, BODY) + lvl(2, UNIT, mar=0.30, bullet=True) + lvl(3, FIG, mar=0.60, bullet=True)
            + "".join(lvl(n, FIG, mar=0.60, bullet=True) for n in range(4, 10)) + "</p:bodyStyle>")
    other = ("<p:otherStyle>" + "".join(
        f'<a:lvl{n}pPr marL="0" algn="l" defTabSz="914400"><a:defRPr sz="1800" kern="0">{fill(scheme="tx1")}'
        '<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/><a:cs typeface="+mn-cs"/></a:defRPr></a:lvl{n}pPr>'.replace("{n}", str(n))
        for n in range(1, 10)) + "</p:otherStyle>")
    return f'<p:txStyles {nsdecls("p", "a")}>{title}{body}{other}</p:txStyles>'


def _replace_children(el, new_xml):
    """Replace the element `el` (in its parent) with the element parsed from new_xml."""
    new = parse_xml(new_xml)
    el.getparent().replace(el, new)
    return new


def build_template(prs, V):
    """Rewrite the master and turn the default template's layouts into the ten Vellum layouts."""
    master = prs.slide_master
    mel = master._element
    csld = mel.find(qn("p:cSld"))
    bg = csld.find(qn("p:bg"))
    new_bg = parse_xml(f'<p:bg {nsdecls("p", "a")}><p:bgPr>{fill(scheme="bg1")}<a:effectLst/></p:bgPr></p:bg>')
    if bg is not None: csld.replace(bg, new_bg)
    else: csld.insert(0, new_bg)
    _replace_children(csld.find(qn("p:spTree")), master_xml(V))
    _replace_children(mel.find(qn("p:txStyles")), tx_styles(V))

    lst = mel.find(qn("p:sldLayoutIdLst"))
    ids = list(lst)
    specs = layouts(V)
    assert len(ids) >= len(specs)
    for el, (name, ltype, show, bgs, S) in zip(ids, specs):
        lp = master.part.related_part(el.get(qn("r:id")))
        root = lp._element
        for k in list(root.attrib):
            if k in ("type", "showMasterSp", "preserve", "userDrawn"): del root.attrib[k]
        if ltype: root.set("type", ltype)
        root.set("preserve", "1")
        if not show: root.set("showMasterSp", "0")
        lcsld = root.find(qn("p:cSld"))
        lcsld.set("name", name)
        old_bg = lcsld.find(qn("p:bg"))
        if old_bg is not None: lcsld.remove(old_bg)
        if bgs:
            lcsld.insert(0, parse_xml(f'<p:bg {nsdecls("p", "a")}><p:bgPr>{fill(scheme=bgs)}<a:effectLst/></p:bgPr></p:bg>'))
        _replace_children(lcsld.find(qn("p:spTree")), S.tree())
    for el in ids[len(specs):]:                    # the default template's extra layouts go
        lst.remove(el)
        master.part.drop_rel(el.get(qn("r:id")))
    return {lay.name: lay for lay in prs.slide_layouts}


# ============================================================================ slide helpers
def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    raise KeyError(idx)


def drop_unused(slide):
    """Remove placeholders left empty (unfilled picture, chart or table placeholders too), so no prompt text shows
    when the deck is edited. The slide number holds a field, not text, and stays."""
    for p in list(slide.placeholders):
        if p.placeholder_format.idx == 12: continue
        unfilled = p.__class__.__name__ in ("PicturePlaceholder", "ChartPlaceholder", "TablePlaceholder")
        if unfilled or (p.has_text_frame and not p.text_frame.text.strip()):
            p._element.getparent().remove(p._element)


def put(p, text, size=None, color=None):
    """Set a placeholder's text, one paragraph per item of a list."""
    tf = p.text_frame
    items = text if isinstance(text, list) else [text]
    tf.text = items[0]
    for t in items[1:]:
        tf.add_paragraph().text = t
    for para in tf.paragraphs:
        for r in para.runs:
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = rgb(color)


def add_footer(slide, layout, text):
    """The footer and the slide number: python-pptx does not copy these placeholders from the layout, so clone them."""
    for idx in (11, 12):
        lph = next(p for p in layout.placeholders if p.placeholder_format.idx == idx)
        slide.shapes.clone_placeholder(lph)
    fits(text, "body", FOOT, FOOT_BOX[2], 1)
    ph(slide, 11).text_frame.text = text
    num_ph = ph(slide, 12)
    p = num_ph.text_frame.paragraphs[0]
    for r in list(p._p):
        p._p.remove(r)
    fld = parse_xml(f'<a:fld {nsdecls("a")} id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">'
                    f'<a:rPr lang="en-GB"/><a:t>{len(slide.part.package.presentation_part.presentation.slides)}</a:t></a:fld>')
    p._p.append(fld)


def claim(slide, text, unit=None, max_lines=2):
    """Set the claim and its unit line; returns the top of the content area under them (inches).
    A one-line claim pulls the unit line and the content up, so the gaps stay the same on every slide."""
    k = fits(text, "title", CLAIM, CW, max_lines)
    t = slide.shapes.title
    t.text = text
    lh = line_h("title", CLAIM, 0.95)
    t.left, t.top, t.width, t.height = Inches(M), Inches(TITLE_Y), Inches(CW), Inches(k * lh + 0.06)
    y = TITLE_Y + k * lh + 0.16
    if unit is not None:
        fits(unit, "body", UNIT, CW, 1)
        u = ph(slide, 13)
        u.text = unit
        u.left, u.top, u.width, u.height = Inches(M), Inches(y), Inches(CW), Inches(0.36)
        y += 0.36
    return y + 0.30


def fit_picture(p, img, box, scale=None, middle=False):
    """Fill a picture placeholder with a PIL image, uncropped, fitted to box = (x, y, w, h), left aligned and at the
    top (or centred in the height with middle=True). `scale` (inches per px) fixes the scale, so that two panels of
    one figure keep the same type size. Returns the picture's bottom edge in inches."""
    buf = io.BytesIO(); img.save(buf, "PNG", optimize=True); buf.seek(0)
    pic = p.insert_picture(buf)
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
    x, y, w, hh = box
    iw, ih = img.size
    s = scale or min(w / iw, hh / ih)
    if middle: y += (hh - ih * s) / 2
    pic.left, pic.top, pic.width, pic.height = Inches(x), Inches(y), Inches(iw * s), Inches(ih * s)
    return y + ih * s


def place(shape, x, y, w, hh):
    """Position a placeholder on the slide: all four values, since one alone writes zeros for the others."""
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(hh)


# ---- figures from the gallery, without their own title block and footer
def bands(img, axis=0, tol=10, gap=6):
    """Runs of rows (axis=0) or columns (axis=1) that hold anything but the background, merged across gaps < gap px."""
    import numpy as np
    a = np.asarray(img.convert("RGB")).astype(int)
    bg = a[2, 2]
    ink = (np.abs(a - bg).max(axis=2) > tol)
    prof = ink.sum(axis=1 - axis) if axis == 0 else ink.sum(axis=0)
    runs, start, last = [], None, None
    for i, v in enumerate(prof):
        if v > 0:
            if start is None: start = i
            elif i - last > gap:
                runs.append((start, last)); start = i
            last = i
    if start is not None: runs.append((start, last))
    return runs


def figure(group, name, mode, skip_top=2, skip_bottom=1, pad=14):
    """A gallery PNG with its title block (title and subtitle bands) and its footer cut off: the slide says both."""
    img = Image.open(os.path.join(FIGURES, group, f"{name}_{mode}.png")).convert("RGB")
    rows = bands(img, 0)
    top = rows[skip_top][0] - pad
    bottom = rows[-1 - skip_bottom][1] + pad
    return img.crop((0, max(0, top), img.width, min(img.height, bottom)))


def split_panels(img, pad=10):
    """Cut a two-panel figure at the widest empty run of columns in its middle half."""
    cols = bands(img, 1, gap=30)
    w = img.width
    best = max(((b[0] - a[1], a[1], b[0]) for a, b in zip(cols, cols[1:]) if 0.25 * w < (a[1] + b[0]) / 2 < 0.75 * w))
    _, l, r = best
    left = img.crop((0, 0, l + pad, img.height))
    right = img.crop((r - pad, 0, w, img.height))
    return [trim(left), trim(right)]


def trim(img, pad=10):
    rows, cols = bands(img, 0), bands(img, 1)
    return img.crop((max(0, cols[0][0] - pad), max(0, rows[0][0] - pad),
                     min(img.width, cols[-1][1] + pad), min(img.height, rows[-1][1] + pad)))


def manifest(group, name):
    import json
    items = json.load(open(os.path.join(FIGURES, group, "manifest.json"), encoding="utf-8"))
    return next(e for e in items if e["id"] == name)


# ============================================================================ native chart styling (lxml where python-pptx has no API)
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _nofill_sppr():
    return parse_xml(f'<c:spPr {nsdecls("c", "a")}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>')


def chart_frame(chart):
    """No border, no fill, square corners, no title, no legend, Corbel throughout."""
    cs = chart._chartSpace
    rc = cs.find(C + "roundedCorners")
    if rc is None:
        rc = parse_xml(f'<c:roundedCorners {nsdecls("c")} val="0"/>')
        anchor = cs.find(C + "date1904")
        if anchor is not None: anchor.addnext(rc)
        else: cs.insert(0, rc)
    rc.set("val", "0")
    old = cs.find(C + "spPr")
    if old is not None: cs.remove(old)
    cs.find(C + "chart").addnext(_nofill_sppr())
    chart.has_title = False
    chart.has_legend = False
    pa = cs.find(C + "chart").find(C + "plotArea")
    old = pa.find(C + "spPr")
    if old is not None: pa.remove(old)
    pa.append(_nofill_sppr())
    chart.font.name = FACES["body"]
    chart.font.size = Pt(FIG)


def plot_layout(chart, x, y, w, hh):
    """Manual inner plot area, as fractions of the chart frame (room on the right for the end labels)."""
    pa = chart._chartSpace.find(C + "chart").find(C + "plotArea")
    lay = pa.find(C + "layout")
    new = parse_xml(f'<c:layout {nsdecls("c")}><c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
                    f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{hh}"/></c:manualLayout></c:layout>')
    if lay is not None: pa.replace(lay, new)
    else: pa.insert(0, new)


def axis_text(axis, V, color, size=FIG):
    f = axis.tick_labels.font
    f.size = Pt(size); f.color.rgb = rgb(color); f.name = FACES["body"]


def value_axis(chart, V, lo, hi, step, fmt):
    va = chart.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = lo, hi, step
    va.has_major_gridlines, va.has_minor_gridlines = True, False
    gl = va.major_gridlines.format.line
    gl.color.rgb = rgb(V.grid); gl.width = Pt(0.75)
    va.format.line.fill.background()
    va.major_tick_mark = va.minor_tick_mark = XL_TICK_MARK.NONE
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = fmt, False
    axis_text(va, V, V.muted)
    return va


def category_axis(chart, V, color, skip=None):
    ca = chart.category_axis
    ca.format.line.color.rgb = rgb(V.grid); ca.format.line.width = Pt(1.0)
    ca.major_tick_mark = ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.has_major_gridlines = ca.has_minor_gridlines = False
    axis_text(ca, V, color)
    if skip:
        el = ca._element
        nm = el.find(C + "noMultiLvlLbl")
        for tag in ("tickLblSkip", "tickMarkSkip"):
            old = el.find(C + tag)
            if old is not None: el.remove(old)
            new = parse_xml(f'<c:{tag} {nsdecls("c")} val="{skip}"/>')
            if nm is not None: nm.addprevious(new)
            else: el.append(new)
    return ca


def round_line(ser):
    """Round caps and joins on a series line, as matplotlib draws them in vellum.py."""
    ln = ser._element.find(C + "spPr").find(A + "ln")
    ln.set("cap", "rnd")
    for old in ln.findall(A + "round"): ln.remove(old)
    fillel = ln.find(A + "solidFill")
    fillel.addnext(parse_xml(f'<a:round {nsdecls("a")}/>'))


def label(point, text, V, color, pos, size=FIG):
    dl = point.data_label
    tf = dl.text_frame
    tf.text = text
    for r in tf.paragraphs[0].runs:
        r.font.size = Pt(size); r.font.color.rgb = rgb(color); r.font.name = FACES["body"]
    dl.position = pos


# ============================================================================ data
def energy():
    groups = [("coal", ["coal_electricity"]), ("oil and gas", ["gas_electricity", "oil_electricity"]),
              ("nuclear", ["nuclear_electricity"]),
              ("hydro and bio", ["hydro_electricity", "biofuel_electricity", "other_renewable_exc_biofuel_electricity"]),
              ("wind", ["wind_electricity"]), ("solar", ["solar_electricity"])]
    rows = {int(r["year"]): r for r in csv.DictReader(open(os.path.join(DATA, "sustainability", "owid-energy.csv"), encoding="utf-8"))
            if r["country"] == "World"}
    twh = {y: [sum(float(rows[y][c] or 0) for c in cs) for _, cs in groups] for y in (2000, 2024, 2025)}
    gen = {y: float(rows[y]["electricity_generation"]) for y in twh}
    for y in twh:
        assert abs(sum(twh[y]) - gen[y]) / gen[y] < 2e-3, "the six groups must add up to world generation"
    share = {y: 100 * sum(twh[y][2:]) / gen[y] for y in twh}
    return dict(groups=[g for g, _ in groups], twh=twh, gen=gen, share=share,
                d_low=sum(twh[2025][2:]) - sum(twh[2024][2:]), d_fossil=sum(twh[2025][:2]) - sum(twh[2024][:2]),
                d_gen=gen[2025] - gen[2024])


def thousands(x):
    return f"{x:,.0f}"


# ============================================================================ the deck
LE_DATA = f"Data: {SOURCE}, processed by Our World in Data, CC BY 4.0"
EN_DATA = "Data: Ember (2026) and Energy Institute (2025), via Our World in Data, CC BY 4.0"


def build(mode):
    V = Palette(mode)
    F = facts(); EN = energy()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    set_theme(prs, V)
    L = build_template(prs, V)
    prs.core_properties.title = "Vellum slides" + ("" if mode == "light" else ", dark")
    prs.core_properties.author = AUTHOR
    prs.core_properties.subject = "Vellum figure style: PowerPoint and Canva template with example slides"
    prs.core_properties.keywords = "Vellum; template; slides"
    total = 13

    def new(layout_name):
        s = prs.slides.add_slide(L[layout_name])
        return s, L[layout_name]

    def tail(s, lay, foot):
        add_footer(s, lay, foot)
        drop_unused(s)

    me = f"{SCRIPT} · git {GIT}"

    # 1. Title
    s, lay = new("Title")
    t = "Every slide in this deck takes its colours and type from one Vellum theme"
    fits(t, "title", COVER, 11.4, 2)
    s.shapes.title.text = t
    ph(s, 1).text = "A PowerPoint template with ten layouts, figures from the Vellum gallery and native charts"
    ph(s, 13).text = f"Vellum example deck · {AUTHOR} · {DATE}"
    tail(s, lay, f"Vellum slides, {mode} · {me}")

    # 2. Section
    s, lay = new("Section")
    ph(s, 13).text = "01"
    s.shapes.title.text = "Figures from the gallery"
    ph(s, 14).text = "PNG twins from matplotlib in picture placeholders; the dark deck takes the dark twins"
    tail(s, lay, f"Vellum slides, {mode} · {me}")

    # 3. Claim + figure: Keeling curve
    for group, name, unit in (("science", "co2_keeling", None), ("science", "life_small_multiples", None)):
        m = manifest(group, name)
        s, lay = new("Claim + figure")
        sub = m["subtitle"].split(" · ")[0]
        y = claim(s, m["title"], sub)
        img = figure(group, name, mode)
        fit_picture(ph(s, 14), img, (M, y, CW, CONTENT_BOTTOM - y), middle=True)
        src = {"co2_keeling": "Data: NOAA GML and Scripps, via vega-datasets 2.8.0",
               "life_small_multiples": "Data: UN WPP (2024), HMD (2024) and others, processed by Our World in Data"}[name]
        tail(s, lay, f"{src} · {os.path.basename(m['script'])} · git {GIT}")

    # 5. Claim + two figures: the landscape and its heat-map twin
    m = manifest("threed", "landscape_surface")
    s, lay = new("Claim + two figures")
    y = claim(s, m["title"], "Illustrative landscape w(z₁, z₂) · the same gradient and limits in both panels")
    a, b = split_panels(figure("threed", "landscape_surface", mode))
    half = (CW - GUTTER) / 2
    room = CONTENT_BOTTOM - 0.46 - 0.12 - y                      # height left for the panels above the captions
    scale = min(min(half / im.width, room / im.height) for im in (a, b))
    bottom = y + max(im.height for im in (a, b)) * scale
    for k, (img, cap) in enumerate(((a, "The shape: orthographic view, elevation 24°, azimuth −58°"),
                                    (b, "The values: the same surface as a heat map"))):
        x = M + k * (half + GUTTER)
        fit_picture(ph(s, 14 + k), img, (x, y, half, room), scale=scale)
        fits(cap, "body", FIG, half, 1)
        c = ph(s, 16 + k); c.text = cap
        place(c, x, bottom + 0.14, half, 0.40)
    tail(s, lay, f"Analytic function landscape(z1, z2) in _common.py, illustrative · landscape_3d.py · git {GIT}")

    # 6. Section
    s, lay = new("Section")
    ph(s, 13).text = "02"
    s.shapes.title.text = "Native charts, numbers and tables"
    ph(s, 14).text = "Drawn by python-pptx in the theme's colours and fonts, and editable in PowerPoint"
    tail(s, lay, f"Vellum slides, {mode} · {me}")

    # 7. Chart: native line chart, three regions, end labels
    regions = ["Asia", "Africa", "Europe"]
    le = {}
    for r in load_rows():
        if r["Entity"] in regions and int(r["Year"]) >= 1950:
            le.setdefault(r["Entity"], {})[int(r["Year"])] = float(r["Life expectancy"])
    years = list(range(1950, F["last"] + 1))
    gain = {r: le[r][years[-1]] - le[r][1950] for r in regions}
    regions.sort(key=lambda r: -gain[r])
    s, lay = new("Chart")
    y = claim(s, f"Since 1950 life expectancy rose by {gain[regions[0]]:.0f} years in {regions[0]}, "
                 f"{gain[regions[1]]:.0f} in {regions[1]} and {gain[regions[2]]:.0f} in {regions[2]}",
              "life expectancy at birth, years")
    cd = CategoryChartData()
    cd.categories = [str(v) for v in years]
    for r in regions:
        cd.add_series(r, [round(le[r][v], 2) for v in years], number_format="0.0")
    box = (M, y - 0.15, CW, CONTENT_BOTTOM - y + 0.15)
    gf = ph(s, 14).insert_chart(XL_CHART_TYPE.LINE, cd)
    gf.left, gf.top, gf.width, gf.height = [Inches(v) for v in box]
    ch = gf.chart
    chart_frame(ch)
    lo = 10 * int(min(min(le[r].values()) for r in regions) // 10)
    value_axis(ch, V, lo, 80, 10, "0")
    category_axis(ch, V, V.muted, skip=10)
    va = ch._chartSpace.find(C + "chart").find(C + "plotArea").find(C + "valAx")
    cb = va.find(C + "crossBetween")
    if cb is not None: cb.set("val", "midCat")
    plot_layout(ch, 0.045, 0.06, 0.80, 0.82)
    for k, ser in enumerate(ch.plots[0].series):
        ser.smooth = False
        ser.format.line.color.rgb = rgb(V.tone[k]); ser.format.line.width = Pt(2.25)
        ser.marker.style = XL_MARKER_STYLE.NONE
        round_line(ser)
        last = ser.points[len(years) - 1]
        last.marker.style = XL_MARKER_STYLE.CIRCLE; last.marker.size = 9
        last.marker.format.fill.solid(); last.marker.format.fill.fore_color.rgb = rgb(V.tone[k])
        last.marker.format.line.color.rgb = rgb(V.surface); last.marker.format.line.width = Pt(1.5)
        label(last, f"{regions[k]} {le[regions[k]][years[-1]]:.1f}", V, V.ink, XL_LABEL_POSITION.RIGHT)
    tail(s, lay, f"{LE_DATA} · {me}")

    # 8. Chart: native column chart, 2000 and 2025, keyed by direct labels on the first pair
    tw = EN["twh"]
    ws = (tw[2000][4] + tw[2000][5], tw[2025][4] + tw[2025][5])
    coal = tw[2025][0] - tw[2000][0]
    assert ws[1] - ws[0] > coal, "the claim: wind and solar grew more than coal did"
    s, lay = new("Chart")
    y = claim(s, f"From 2000 to 2025 wind and solar added {thousands(ws[1] - ws[0])} TWh, "
                 f"more than coal's {thousands(coal)} TWh", "world electricity generation by source, TWh")
    cd = CategoryChartData()
    cd.categories = EN["groups"]
    for yr in (2000, 2025):
        cd.add_series(str(yr), [round(v) for v in tw[yr]], number_format="#,##0")
    box = (M, y - 0.15, CW, CONTENT_BOTTOM - y + 0.15)
    gf = ph(s, 14).insert_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, cd)
    gf.left, gf.top, gf.width, gf.height = [Inches(v) for v in box]
    ch = gf.chart
    chart_frame(ch)
    top = 2000 * -(-max(max(v) for v in tw.values()) // 2000)
    value_axis(ch, V, 0, top, 2000, "#,##0")
    category_axis(ch, V, V.secondary)
    plot_layout(ch, 0.06, 0.06, 0.92, 0.82)
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 60, -6
    for k, ser in enumerate(plot.series):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(V.tone[k])
        ser.format.line.fill.background()
        ser.invert_if_negative = False
        label(ser.points[0], ser.name, V, V.secondary, XL_LABEL_POSITION.OUTSIDE_END)
    tail(s, lay, f"{EN_DATA} · {me}")

    # 9. Big number
    s, lay = new("Big number")
    claim(s, f"In 2025 low-carbon sources made {EN['share'][2025]:.1f} % of world electricity", max_lines=1)
    ph(s, 13).text = f"{EN['share'][2025]:.1f} %"
    what = "of world electricity generation in 2025"
    k = fits(what, "body", BODY, 5.7, 2)
    ph(s, 14).text = what
    cmp_ = ph(s, 15)
    cmp_.text = f"up from {EN['share'][2000]:.1f} % in 2000"
    place(cmp_, M + 0.45, BIG_Y + 2.06 + k * line_h("body", BODY, 0.95) + 0.12, 5.7, 0.5)
    assert EN["d_fossil"] < 0 < EN["d_gen"] < EN["d_low"]
    sentence = (f"In 2025 low-carbon sources added {thousands(EN['d_low'])} TWh, more than the {thousands(EN['d_gen'])} TWh "
                f"by which generation grew, so fossil generation fell by {thousands(-EN['d_fossil'])} TWh.")
    fits(sentence, "body", 24, CW - 7.2, 5)
    ph(s, 16).text = sentence
    tail(s, lay, f"{EN_DATA} · {me}")

    # 10. Stat tiles on the claim-only layout
    s, lay = new("Claim only")
    y = claim(s, f"The life-expectancy series covers {F['n']} countries and territories over {F['n_years']} years",
              f"Period life expectancy at birth, {F['start']}–{F['last']}")
    tiles = [(str(F["n"]), "countries and territories", f"{F['country_years']:,} country-years in all"),
             (str(F["n_years"]), "years of annual estimates", f"{F['start']} to {F['last']}, one value a year"),
             (f"{num(-F['drop'])}", f"years, the world's fall from {F['base']} to {F['low']}",
              f"the deepest since {F['since'] - 1}–{str(F['since'])[2:]}")]
    tw_ = (CW - 2 * GUTTER) / 3
    for k, (v, l1, l2) in enumerate(tiles):
        x = M + k * (tw_ + GUTTER)
        th = 2.75
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y + 0.1), Inches(tw_), Inches(th))
        sh.name = f"Stat tile {k + 1}"
        sh.adjustments[0] = 0.12 / th
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(V.card)
        sh.line.fill.background()
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.35); tf.margin_top = Inches(0.30); tf.margin_bottom = Inches(0.2)
        for j, (txt, pt, col, face) in enumerate(((v, STAT, V.ink, FACES["title"]), (l1, UNIT, V.muted, FACES["body"]),
                                                  (l2, UNIT, V.secondary, FACES["body"]))):
            fits(txt, "title" if j == 0 else "body", pt, tw_ - 0.7, 1 if j == 0 else 2)
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(0 if j == 0 else 6)
            r = p.add_run(); r.text = txt
            r.font.size = Pt(pt); r.font.color.rgb = rgb(col); r.font.name = face
    tail(s, lay, f"{LE_DATA} · {me}")

    # 11. Table
    s, lay = new("Table")
    worst, mild = F["worst"], F["mildest"]
    y = claim(s, f"From {F['base']} to {F['low']} life expectancy fell by {num(F['rfall'][mild])} years in {mild} "
                 f"and by {num(F['rfall'][worst])} in {worst}",
              f"Period life expectancy at birth by UN region, years · sorted by the fall")
    rows = sorted(REGIONS, key=lambda r: -F["rfall"][r])
    R = F["regions"]
    header = ["UN region", str(F["base"]), str(F["low"]), str(F["last"]), f"fall, {F['base']}–{str(F['low'])[2:]}"]
    body = [[r, f"{R[r][F['base']]:.1f}", f"{R[r][F['low']]:.1f}", f"{R[r][F['last']]:.1f}", num(-F["rfall"][r])] for r in rows]
    widths = [4.9, 1.6, 1.6, 1.6, 2.3]
    gf = ph(s, 14).insert_table(len(body) + 1, len(header))
    gf.left, gf.top = Inches(M), Inches(y)
    tbl = gf.table
    style_table(tbl, V, [header] + body, widths, highlight=rows.index(worst) + 1)
    gf.width = Inches(sum(widths))
    tail(s, lay, f"{LE_DATA} · {me}")

    # 12. Closing
    s, lay = new("Closing")
    s.shapes.title.text = "Every figure, chart and number here is drawn by a script from open data"
    ph(s, 13).text = f"{AUTHOR} · {URL}"
    put(ph(s, 14), [f"Deck: examples/python/office/{SCRIPT} · figures: examples/python · data: data/SOURCES.md",
                    f"Downloads: vellum-slides-light.pptx, vellum-slides-dark.pptx, vellum-office.tsv · git {GIT}"])
    tail(s, lay, f"Vellum slides, {mode} · {me}")

    # 13. Using this template in PowerPoint and Canva
    s, lay = new("Claim + text")
    y = claim(s, "Using this template in PowerPoint and Canva")
    N = SPEC[mode]["neutrals"]
    tones = " ".join(h(c) for c in V.tone)
    twin = "" if mode == "light" else ", by role in the dark twin"
    notes = [
        "PowerPoint: Home → New Slide lists the ten Vellum layouts; Design → Variants holds the theme. "
        "File → Save As → PowerPoint Template (.potx) keeps it as a template.",
        f"Theme colours{twin}: Dark 1 ink {h(N['ink'])}, Light 1 surface {h(N['surface'])}, Dark 2 secondary "
        f"{h(N['secondary'])}, Light 2 card {h(N['card'])}; Accents 1–6, the tones in slot order: {tones}.",
        "Fonts: Palatino Linotype for headings, Corbel for body. Canva has neither by default: upload them to the "
        "Brand Kit if your licence allows, or set a similar serif and sans such as EB Garamond and Lato.",
        "Canva: drag the .pptx into Canva or use Import file. Text, shapes, pictures and tables stay editable; "
        "check the two native charts, which may arrive as Canva charts or as pictures, and the slide numbers.",
        f"Text stays ink, secondary or muted {h(N['muted'])}; the tones are for data. Every hex value of both twins "
        "is in vellum-office.tsv, ready for the Brand Kit's colours.",
    ]
    body_ph = ph(s, 1)
    place(body_ph, M, y - 0.1, 11.0, CONTENT_BOTTOM - y + 0.1)
    total_lines = sum(fits(n_, "body", 18, 11.0, 3) for n_ in notes)
    assert total_lines <= 12, total_lines
    put(body_ph, notes, size=18)
    tail(s, lay, f"Colours: vellum-office.tsv · fonts: vellum.json · {me}")

    assert len(prs.slides) == total, len(prs.slides)
    for slide in prs.slides:                      # typographic apostrophes in every text run
        for t in slide._element.iter(qn("a:t")):
            if t.text and "'" in t.text: t.text = t.text.replace("'", "’")
    out = os.path.join(DOWNLOADS, f"vellum-slides-{mode}.pptx")
    prs.save(out)
    print("wrote", out)
    return out


def load_rows():
    return csv.DictReader(open(os.path.join(DATA, "owid-life-expectancy.csv"), encoding="utf-8"))


def style_table(tbl, V, cells, widths, highlight=None):
    """Vellum table: ink header on the card, alternate rows on the card, the highlighted row on the terracotta tint,
    a rule-grey hairline under the header and none elsewhere, numbers right-aligned in 18 pt Corbel."""
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1"); tblPr.set("bandRow", "0")
    sid = tblPr.find(A + "tableStyleId")
    if sid is None:
        sid = etree.SubElement(tblPr, A + "tableStyleId")
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"          # No Style, No Grid: the cells carry everything
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(cells):
        tbl.rows[i].height = Inches(0.56 if i == 0 else 0.52)
        for j, txt in enumerate(row):
            cell = tbl.cell(i, j)
            bgc = (V.card if i == 0 or i % 2 == 0 else V.surface)
            if highlight is not None and i == highlight:
                bgc = V.tint[SPEC["order"].index("terracotta")]
            tcPr = cell._tc.get_or_add_tcPr()
            for child in list(tcPr): tcPr.remove(child)
            below = (f'<a:lnB w="{int(Pt(1.0))}"><a:solidFill><a:srgbClr val="{h(V.rule)}"/></a:solidFill></a:lnB>' if i == 0
                     else '<a:lnB w="0"><a:noFill/></a:lnB>')
            for x in ('<a:lnL w="0"><a:noFill/></a:lnL>', '<a:lnR w="0"><a:noFill/></a:lnR>', '<a:lnT w="0"><a:noFill/></a:lnT>', below,
                      f'<a:solidFill><a:srgbClr val="{h(bgc)}"/></a:solidFill>'):
                tcPr.append(parse_xml(x.replace(">", f' {nsdecls("a")}>', 1)))
            cell.margin_left = cell.margin_right = Inches(0.18)
            cell.margin_top = cell.margin_bottom = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.text = txt
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            for r in p.runs:
                r.font.size = Pt(18); r.font.name = FACES["body"]
                r.font.color.rgb = rgb(V.ink if (i == 0 or j == 0 or j == len(row) - 1) else V.secondary)
                r.font.bold = i == 0


# ============================================================================ previews through PowerPoint
PS = r"""
param([string]$pptx, [string]$outdir, [string]$pdf)
$ErrorActionPreference = 'Stop'
$pp = New-Object -ComObject PowerPoint.Application
try {
  $pp.DisplayAlerts = 1
  $pres = $pp.Presentations.Open($pptx, -1, 0, 0)
  $n = $pres.Slides.Count
  for ($i = 1; $i -le $n; $i++) {
    $pres.Slides.Item($i).Export((Join-Path $outdir ('slide-{0:D2}.png' -f $i)), 'PNG', 1600, 900)
  }
  if ($pdf) { $pres.SaveAs($pdf, 32) }
  $pres.Close()
  Write-Output "opened and exported $n slides"
} finally {
  $pp.Quit()
  [System.Runtime.InteropServices.Marshal]::ReleaseComObject($pp) | Out-Null
  [GC]::Collect(); [GC]::WaitForPendingFinalizers()
}
"""


def render(pptx, mode):
    out = os.path.join(MEDIA, mode)
    os.makedirs(out, exist_ok=True)
    for f in os.listdir(out):
        if f.startswith("slide-") and f.endswith(".png"): os.remove(os.path.join(out, f))
    pdf = os.path.join(MEDIA, f"vellum-slides-{mode}.pdf") if mode == "light" else ""
    with tempfile.TemporaryDirectory() as tmp:
        ps1 = os.path.join(tmp, "render.ps1")
        open(ps1, "w", encoding="utf-8").write(PS)
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", ps1,
                            "-pptx", os.path.abspath(pptx), "-outdir", os.path.abspath(out), "-pdf", pdf],
                           capture_output=True, text=True, timeout=600)
    print(r.stdout.strip(), r.stderr.strip())
    if r.returncode != 0:
        raise RuntimeError(f"PowerPoint could not open or export {pptx}")
    pages = sorted(f for f in os.listdir(out) if f.startswith("slide-"))
    return [f"media/slides/{mode}/{f}" for f in pages]


if __name__ == "__main__":
    os.makedirs(DOWNLOADS, exist_ok=True)
    files = {mode: build(mode) for mode in ("light", "dark")}
    if "--no-render" in sys.argv:
        sys.exit(0)
    pages = {mode: render(files[mode], mode) for mode in ("light", "dark")}
    record("slides_pptx", "pptx", dict(
        title="Every slide in this deck takes its colours and type from one Vellum theme",
        subtitle="A PowerPoint template that Canva can import: Vellum theme colours and fonts, ten layouts, thirteen example slides, light and dark",
        kind="slides", form="PowerPoint template, 13 slides", medium="slides", tool="PowerPoint (python-pptx)",
        data=("Figures from the Vellum gallery; UN World Population Prospects (2024) and Ember (2026) with the Energy "
              "Institute (2025), processed by Our World in Data, CC BY 4.0"),
        script=rel(__file__),
        alt=("A 16:9 PowerPoint deck in the Vellum style, light and dark: a title slide, two section slides on the card, "
             "the Keeling curve and the life-expectancy small multiples as claim-and-figure slides, the adaptive "
             "landscape and its heat-map twin side by side, a native line chart of life expectancy in Asia, Africa and "
             "Europe with end labels, a native column chart of world electricity by source in 2000 and 2025, a big "
             "number (42.6 % low-carbon electricity in 2025), three stat tiles, a regional table, a closing slide and a "
             "slide on using the template in PowerPoint and Canva."),
        note=("Theme colours, fonts, master and layouts are set in the .pptx itself; previews exported by PowerPoint. "
              "Downloads: downloads/vellum-slides-light.pptx and downloads/vellum-slides-dark.pptx."),
        href="downloads/vellum-slides-light.pptx",
        # the gallery cover is slide 7, the native line chart
        files={"light": pages["light"][6], "dark": pages["dark"][6], "pages": pages["light"],
               "pages_dark": pages["dark"], "pdf": "media/slides/vellum-slides-light.pdf",
               "pptx": "downloads/vellum-slides-light.pptx", "pptx_dark": "downloads/vellum-slides-dark.pptx"}))
    print("recorded slides_pptx/pptx")
